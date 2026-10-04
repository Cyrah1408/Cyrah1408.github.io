#!/usr/bin/env python3
"""Briefly data tool: converts between the site's edition JSON and data/briefly.xlsx.

The website reads data/briefly.xlsx directly. The daily workflow should never edit
spreadsheet cells by hand; it works on the familiar JSON and converts:

  python3 tools/briefly_data.py export data/briefly.xlsx editions.json   # xlsx -> JSON
  python3 tools/briefly_data.py import editions.json data/briefly.xlsx   # JSON -> xlsx (+ data/version.json)
  python3 tools/briefly_data.py validate data/briefly.xlsx                # checks only

JSON shape: {"editions":[ ...newest first... ]} (same as the old embedded briefly-data block).
`import` validates first and refuses to write anything if a check fails.
"""
import json, os, re, sys, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

SCHEMA_VERSION = 1
MAX_EDITIONS = 7
SEP = " | "          # separator for short lists inside one cell
THEMES = {"geopolitics", "indiapol", "environment", "science", "tech", "economy",
          "literature", "history", "arts", "sport"}

# Sheet name -> column headers. This is the schema; keep it identical in assets/briefly-loader.js.
SHEETS = {
    "editions": ["date", "no", "label", "greeting", "para_1", "para_2", "para_3",
                 "hive_letters", "hive_centre", "hive_words", "timeline_prompt",
                 "debate_motion", "debate_for_1", "debate_for_2", "debate_against_1",
                 "debate_against_2", "debate_opinion"],
    "stories": ["edition_date", "section", "seq", "id", "theme", "head", "summary", "why", "cls",
                "stat_big", "stat_label", "minutes", "dek", "body_1", "body_2", "body_3", "body_4",
                "note", "checks", "think", "game_type", "game_word", "game_hint"],
    "sources": ["edition_date", "item_id", "seq", "outlet", "url"],
    "pairs": ["edition_date", "parent_id", "seq", "term", "definition"],
    "opinions": ["edition_date", "seq", "id", "theme", "outlet", "author", "role", "date", "head",
                 "url", "gist", "other", "think"],
    "games": ["edition_date", "seq", "id", "type", "word", "hint"],
    "connections": ["edition_date", "seq", "name", "words"],
    "timeline": ["edition_date", "seq", "event", "year"],
    "duel": ["edition_date", "seq", "a_label", "a_value", "b_label", "b_value", "note"],
    "audio": ["edition_date", "track", "seq", "speaker", "text", "tone"],
    "meta": ["key", "value"],
}


class DataError(Exception):
    pass


def plist(items, where):
    for x in items:
        if "|" in str(x):
            raise DataError(f"{where}: list item contains '|', which is reserved: {x!r}")
    return SEP.join(str(x) for x in items)


def split(v):
    v = s(v)
    return [x.strip() for x in v.split("|")] if v.strip() else []


def s(v):
    return "" if v is None else str(v)


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    f = float(v)
    return int(f) if f.is_integer() else f


def fixed(lst, n, where):
    lst = list(lst or [])
    if len(lst) > n:
        raise DataError(f"{where}: has {len(lst)} entries; the sheet allows at most {n}")
    return lst + [""] * (n - len(lst))


# ---------------------------------------------------------------- JSON -> rows
def to_rows(data):
    eds = data["editions"]
    R = {k: [] for k in SHEETS}
    for e in eds:
        d = e["date"]
        b = e.get("bulletin", {})
        hv, tl, db = e.get("hive", {}), e.get("timeline", {}), e.get("debate", {})
        paras = fixed(b.get("paras"), 3, f"{d} bulletin.paras")
        f2, a2 = fixed(db.get("for"), 2, f"{d} debate.for"), fixed(db.get("against"), 2, f"{d} debate.against")
        R["editions"].append([d, e["no"], e["label"], b.get("greeting", ""), *paras,
                              hv.get("letters", ""), hv.get("centre", ""), plist(hv.get("words", []), f"{d} hive"),
                              tl.get("prompt", ""), db.get("motion", ""), *f2, *a2, db.get("opinion", "")])

        def src(item):
            for i, (o, u) in enumerate(item.get("sources") or [], 1):
                R["sources"].append([d, item["id"], i, o, u])

        for sec in ("brief", "curated", "desk"):
            for i, it in enumerate(e.get(sec) or [], 1):
                st, g = it.get("stat") or {}, it.get("game") or {}
                body = fixed(it.get("body"), 4, f"{it['id']} body")
                R["stories"].append([d, sec, i, it["id"], it["theme"], it.get("head", ""), it.get("summary", ""),
                                     it.get("why", ""), it.get("cls", ""), st.get("big", ""), st.get("label", ""),
                                     it.get("minutes"), it.get("dek", ""), *body, it.get("note", ""),
                                     plist(it.get("checks", []), it["id"]), plist(it.get("think", []), it["id"]),
                                     g.get("type", ""), g.get("word", ""), g.get("hint", "")])
                src(it)
                for j, (t, df) in enumerate(it.get("words") or [], 1):
                    R["pairs"].append([d, it["id"], j, t, df])
        for i, o in enumerate(e.get("opinions") or [], 1):
            R["opinions"].append([d, i, o["id"], o["theme"], o["outlet"], o["author"], o.get("role", ""),
                                  o.get("date", ""), o["head"], o["url"], o.get("gist", ""), o.get("other", ""),
                                  plist(o.get("think", []), o["id"])])
        for i, g in enumerate(e.get("games") or [], 1):
            R["games"].append([d, i, g["id"], g["type"], g.get("word", ""), g.get("hint", "")])
            for j, (t, df) in enumerate(g.get("pairs") or [], 1):
                R["pairs"].append([d, g["id"], j, t, df])
        for i, gr in enumerate((e.get("connections") or {}).get("groups") or [], 1):
            R["connections"].append([d, i, gr["name"], plist(gr["words"], f"{d} connections")])
        for i, ev in enumerate(tl.get("events") or [], 1):
            R["timeline"].append([d, i, ev["t"], ev["y"]])
        for i, r in enumerate(e.get("duel") or [], 1):
            R["duel"].append([d, i, r["a"][0], r["a"][1], r["b"][0], r["b"][1], r.get("note", "")])
        for track, v in (e.get("audio") or {}).items():
            for i, ln in enumerate(v.get("lines") or [], 1):
                R["audio"].append([d, track, i, ln[0], ln[1], ln[2] if len(ln) > 2 else ""])
    return R


# ---------------------------------------------------------------- rows -> JSON
def from_rows(R):
    """R: sheet -> list of dicts. Must mirror buildEditions() in assets/briefly-loader.js."""
    def by(sheet, key="edition_date"):
        out = {}
        for r in R.get(sheet, []):
            out.setdefault(s(r.get(key)), []).append(r)
        for v in out.values():
            v.sort(key=lambda r: num(r.get("seq")) or 0)
        return out

    stories, sources, pairs = by("stories"), by("sources", "item_id"), by("pairs", "parent_id")
    opinions, games, conns = by("opinions"), by("games"), by("connections")
    tls, duels, audio = by("timeline"), by("duel"), by("audio")
    srcs = lambda iid: [[s(x["outlet"]), s(x["url"])] for x in sources.get(iid, [])]
    prs = lambda iid: [[s(x["term"]), s(x["definition"])] for x in pairs.get(iid, [])]

    eds = []
    for r in R["editions"]:
        d = s(r["date"])
        e = {"date": d, "no": num(r["no"]), "label": s(r["label"])}
        au = {}
        for a in audio.get(d, []):
            au.setdefault(s(a["track"]), {"lines": []})["lines"].append([s(a["speaker"]), s(a["text"]), s(a["tone"])])
        if au:
            e["audio"] = au
        e["bulletin"] = {"greeting": s(r["greeting"]),
                         "paras": [s(r[k]) for k in ("para_1", "para_2", "para_3") if s(r[k])]}
        sec = {"brief": [], "curated": [], "desk": []}
        for x in stories.get(d, []):
            iid, kind = s(x["id"]), s(x["section"])
            if kind == "brief":
                it = {"id": iid, "theme": s(x["theme"]), "stat": {"big": s(x["stat_big"]), "label": s(x["stat_label"])},
                      "head": s(x["head"]), "summary": s(x["summary"]), "why": s(x["why"]), "cls": s(x["cls"]),
                      "words": prs(iid), "sources": srcs(iid)}
            elif kind == "curated":
                it = {"id": iid, "theme": s(x["theme"]), "minutes": num(x["minutes"]), "head": s(x["head"]),
                      "dek": s(x["dek"]), "body": [s(x[f"body_{i}"]) for i in range(1, 5) if s(x[f"body_{i}"])],
                      "game": {"type": s(x["game_type"]), "word": s(x["game_word"]), "hint": s(x["game_hint"])},
                      "note": s(x["note"]), "checks": split(x["checks"]), "think": split(x["think"]),
                      "sources": srcs(iid)}
            else:
                it = {"id": iid, "theme": s(x["theme"]), "head": s(x["head"]), "summary": s(x["summary"]),
                      "why": s(x["why"])}
                if s(x["stat_big"]) or s(x["stat_label"]):
                    it["stat"] = {"big": s(x["stat_big"]), "label": s(x["stat_label"])}
                it["sources"] = srcs(iid)
            sec.setdefault(kind, []).append(it)
        e["brief"] = sec["brief"]
        e["games"] = []
        for g in games.get(d, []):
            gi = {"id": s(g["id"]), "type": s(g["type"])}
            if s(g["word"]) or s(g["hint"]):
                gi["word"], gi["hint"] = s(g["word"]), s(g["hint"])
            if pairs.get(gi["id"]):
                gi["pairs"] = prs(gi["id"])
            e["games"].append(gi)
        e["connections"] = {"groups": [{"name": s(c["name"]), "words": split(c["words"])} for c in conns.get(d, [])]}
        e["hive"] = {"letters": s(r["hive_letters"]), "centre": s(r["hive_centre"]), "words": split(r["hive_words"])}
        e["timeline"] = {"prompt": s(r["timeline_prompt"]),
                         "events": [{"t": s(t["event"]), "y": num(t["year"])} for t in tls.get(d, [])]}
        e["duel"] = [{"a": [s(x["a_label"]), num(x["a_value"])], "b": [s(x["b_label"]), num(x["b_value"])],
                      "note": s(x["note"])} for x in duels.get(d, [])]
        e["debate"] = {"motion": s(r["debate_motion"]),
                       "for": [s(r[k]) for k in ("debate_for_1", "debate_for_2") if s(r[k])],
                       "against": [s(r[k]) for k in ("debate_against_1", "debate_against_2") if s(r[k])],
                       "opinion": s(r["debate_opinion"])}
        e["opinions"] = [{"id": s(o["id"]), "theme": s(o["theme"]), "outlet": s(o["outlet"]), "author": s(o["author"]),
                          "role": s(o["role"]), "date": s(o["date"]), "head": s(o["head"]), "url": s(o["url"]),
                          "gist": s(o["gist"]), "other": s(o["other"]), "think": split(o["think"])}
                         for o in opinions.get(d, [])]
        e["curated"] = sec["curated"]
        if sec["desk"]:
            e["desk"] = sec["desk"]
        eds.append(e)
    eds.sort(key=lambda e: e["date"], reverse=True)
    return {"editions": eds}


# ---------------------------------------------------------------- checks
def validate(data):
    errs = []
    eds = data.get("editions")
    if not isinstance(eds, list) or not eds:
        raise DataError("no editions")
    dates = [e.get("date", "") for e in eds]
    if len(set(dates)) != len(dates):
        errs.append("duplicate edition dates")
    if dates != sorted(dates, reverse=True):
        errs.append("editions are not newest first")
    ids = set()
    for e in eds:
        d = e.get("date", "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", d):
            errs.append(f"bad edition date {d!r}")
        for k in ("no", "label", "bulletin", "brief", "curated"):
            if k not in e:
                errs.append(f"{d}: missing {k}")
        items = [*e.get("brief", []), *e.get("curated", []), *(e.get("desk") or []), *(e.get("opinions") or []),
                 *(e.get("games") or [])]
        for it in items:
            iid = it.get("id")
            if not iid:
                errs.append(f"{d}: item without id")
            elif iid in ids:
                errs.append(f"duplicate id {iid}")
            ids.add(iid)
            if "theme" in it and it["theme"] not in THEMES:
                errs.append(f"{iid}: unknown theme {it['theme']!r}")
            for o, u in it.get("sources") or []:
                if not str(u).startswith("http"):
                    errs.append(f"{iid}: bad source url {u!r}")
        for o in e.get("opinions") or []:
            if not str(o.get("url", "")).startswith("http"):
                errs.append(f"{o.get('id')}: bad url")
        for ln in [l for a in (e.get("audio") or {}).values() for l in a.get("lines", [])]:
            if ln[0] not in ("F", "M"):
                errs.append(f"{d}: audio speaker must be F or M, got {ln[0]!r}")
    if errs:
        raise DataError("; ".join(errs[:20]) + (f" (+{len(errs)-20} more)" if len(errs) > 20 else ""))


def canon(x):
    return json.dumps(x, sort_keys=True, ensure_ascii=False)


# ---------------------------------------------------------------- xlsx I/O
def write_xlsx(data, path):
    validate(data)
    data = {"editions": data["editions"][:MAX_EDITIONS]}
    R = to_rows(data)
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    R["meta"] = [["schema_version", str(SCHEMA_VERSION)], ["updated_at", now],
                 ["latest_edition", data["editions"][0]["date"]], ["editions", str(len(data["editions"]))]]
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for row in R[name]:
            ws.append([None if v == "" else v for v in row])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = 14 if c in ("edition_date", "seq", "date", "no", "section", "theme", "speaker", "tone", "year") else 40
        if R[name]:
            t = Table(displayName=f"t_{name}", ref=f"A1:{get_column_letter(len(cols))}{len(R[name]) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(t)
    tmp = path[:-5] + ".tmp.xlsx"
    wb.save(tmp)
    # Round-trip proof before replacing the live file.
    back = read_xlsx(tmp)
    if canon(back) != canon(data):
        os.remove(tmp)
        raise DataError("round-trip check failed: the spreadsheet would not reproduce the JSON exactly")
    os.replace(tmp, path)
    vpath = os.path.join(os.path.dirname(os.path.abspath(path)), "version.json")
    with open(vpath, "w") as f:
        json.dump({"version": now, "latest": data["editions"][0]["date"],
                   "label": data["editions"][0]["label"], "editions": len(data["editions"]),
                   "schema": SCHEMA_VERSION}, f)
        f.write("\n")
    return now


def read_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    R = {}
    for name, cols in SHEETS.items():
        if name not in wb.sheetnames:
            raise DataError(f"sheet {name!r} is missing")
        rows = list(wb[name].iter_rows(values_only=True))
        head = [s(h) for h in (rows[0] if rows else [])]
        if head[:len(cols)] != cols:
            raise DataError(f"sheet {name!r} columns changed: expected {cols}, found {head}")
        R[name] = [dict(zip(cols, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    wb.close()
    return from_rows(R)


def main(argv):
    if len(argv) < 3:
        print(__doc__)
        return 2
    cmd = argv[1]
    try:
        if cmd == "export":
            data = read_xlsx(argv[2])
            validate(data)
            with open(argv[3], "w") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
            e = data["editions"][0]
            print(f"exported {len(data['editions'])} editions; newest {e['date']} ({e['label']})")
        elif cmd == "import":
            with open(argv[2]) as f:
                data = json.load(f)
            v = write_xlsx(data, argv[3])
            e = data["editions"][0]
            print(f"wrote {argv[3]} ({os.path.getsize(argv[3])} bytes), version {v}; newest {e['date']} ({e['label']})")
        elif cmd == "validate":
            data = read_xlsx(argv[2])
            validate(data)
            print(f"ok: {len(data['editions'])} editions; newest {data['editions'][0]['date']}")
        else:
            print(__doc__)
            return 2
    except DataError as ex:
        print(f"ERROR: {ex}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
